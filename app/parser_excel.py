from pathlib import Path
from openpyxl import load_workbook
from .models import Record, records_from_rows


def parse(path: Path) -> list[Record]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [[str(c) if c is not None else '' for c in row]
                for row in workbook.active.iter_rows(values_only=True)]
        return records_from_rows(rows)
    finally:
        workbook.close()
